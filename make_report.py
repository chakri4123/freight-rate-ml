"""Render REPORT.md (+ the scorer chart) to REPORT.docx. Run after train.py and score.py."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parent


def add_inline(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        elif part:
            paragraph.add_run(part)


def main() -> None:
    doc = Document()
    lines = (ROOT / "REPORT.md").read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r, cells in enumerate(rows):
                for c, text in enumerate(cells):
                    add_inline(table.cell(r, c).paragraphs[0], text)
            continue
        m = re.match(r"!\[.*\]\((.+)\)", line)
        if m:
            doc.add_picture(str(ROOT / m.group(1)), width=Inches(6.5))
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("#").strip(), level=level - 1 if level > 1 else 0)
        elif re.match(r"^\s*(-|\d+\.) ", line):
            style = "List Number" if re.match(r"^\s*\d+\.", line) else "List Bullet"
            add_inline(doc.add_paragraph(style=style), re.sub(r"^\s*(-|\d+\.) ", "", line))
        elif line.strip():
            add_inline(doc.add_paragraph(), line)
        i += 1
    doc.save(ROOT / "REPORT.docx")
    print("Wrote REPORT.docx")


if __name__ == "__main__":
    main()
